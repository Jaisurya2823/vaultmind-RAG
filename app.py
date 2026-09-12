#!/usr/bin/env python3
"""
VaultMind v2 — Flask + SQLite + ChromaDB + Groq RAG
Run: python app.py  →  http://localhost:5000
"""

import os, json, shutil, threading
from pathlib import Path
from flask import Flask, request, jsonify, send_from_directory
from dotenv import load_dotenv

load_dotenv()

app = Flask(__name__, static_folder="static")

# ─── Config ───────────────────────────────────────────────────────────────────
DB_PATH       = "vaultmind.db"
CHROMA_BASE   = "./chroma_db"
UPLOADS_DIR   = "./uploads"
GROQ_KEY      = os.getenv("GROQ_API_KEY", "")
GROQ_MODEL    = "llama3-8b-8192"
EMBED_MODEL   = "all-MiniLM-L6-v2"
CHUNK_SIZE    = 1000
CHUNK_OVERLAP = 150
TOP_K         = 4

os.makedirs(UPLOADS_DIR, exist_ok=True)
os.makedirs(CHROMA_BASE,  exist_ok=True)

# ─── Database ─────────────────────────────────────────────────────────────────
import sqlite3

def db():
    conn = sqlite3.connect(DB_PATH, check_same_thread=False)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn

def init_db():
    with db() as c:
        c.executescript("""
        CREATE TABLE IF NOT EXISTS projects (
            id           INTEGER PRIMARY KEY AUTOINCREMENT,
            name         TEXT    NOT NULL,
            description  TEXT    DEFAULT '',
            instructions TEXT    DEFAULT '',
            created_at   TEXT    DEFAULT (datetime('now')),
            updated_at   TEXT    DEFAULT (datetime('now'))
        );
        CREATE TABLE IF NOT EXISTS conversations (
            id           INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id   INTEGER REFERENCES projects(id) ON DELETE CASCADE,
            title        TEXT    NOT NULL DEFAULT 'New conversation',
            created_at   TEXT    DEFAULT (datetime('now')),
            updated_at   TEXT    DEFAULT (datetime('now'))
        );
        CREATE TABLE IF NOT EXISTS messages (
            id               INTEGER PRIMARY KEY AUTOINCREMENT,
            conversation_id  INTEGER NOT NULL REFERENCES conversations(id) ON DELETE CASCADE,
            role             TEXT    NOT NULL,
            content          TEXT    NOT NULL,
            sources          TEXT    DEFAULT '[]',
            confidence       TEXT    DEFAULT '',
            created_at       TEXT    DEFAULT (datetime('now'))
        );
        CREATE TABLE IF NOT EXISTS project_files (
            id           INTEGER PRIMARY KEY AUTOINCREMENT,
            project_id   INTEGER NOT NULL REFERENCES projects(id) ON DELETE CASCADE,
            filename     TEXT    NOT NULL,
            filepath     TEXT    NOT NULL,
            file_size    INTEGER DEFAULT 0,
            indexed      INTEGER DEFAULT 0,
            created_at   TEXT    DEFAULT (datetime('now'))
        );
        CREATE TABLE IF NOT EXISTS conversation_files (
            id               INTEGER PRIMARY KEY AUTOINCREMENT,
            conversation_id  INTEGER NOT NULL REFERENCES conversations(id) ON DELETE CASCADE,
            filename         TEXT    NOT NULL,
            filepath         TEXT    NOT NULL,
            file_size        INTEGER DEFAULT 0,
            indexed          INTEGER DEFAULT 0,
            created_at       TEXT    DEFAULT (datetime('now'))
        );
        """)

# ─── RAG Engine ───────────────────────────────────────────────────────────────

_emb      = None
_emb_lock = threading.Lock()

def get_emb():
    global _emb
    if _emb is None:
        with _emb_lock:
            if _emb is None:
                from langchain_huggingface import HuggingFaceEmbeddings
                print("[VaultMind] Loading embedding model…")
                _emb = HuggingFaceEmbeddings(
                    model_name=EMBED_MODEL,
                    model_kwargs={"device": "cpu"},
                    encode_kwargs={"normalize_embeddings": True},
                )
                print("[VaultMind] Embeddings ready.")
    return _emb

# ── Supported extension sets ───────────────────────────────────────────────────
_TEXT_EXTS = {
    # Plain text / markup
    ".txt", ".md", ".markdown", ".rst", ".log", ".ini", ".cfg", ".conf",
    ".yaml", ".yml", ".toml", ".env", ".properties", ".editorconfig",
    # Shell / scripting
    ".sh", ".bash", ".zsh", ".fish", ".bat", ".cmd", ".ps1",
    # Python
    ".py", ".pyw", ".pyi", ".ipynb",  # ipynb also has dedicated loader
    # JavaScript / TypeScript
    ".js", ".mjs", ".cjs", ".ts", ".tsx", ".jsx",
    # Web
    ".css", ".scss", ".sass", ".less", ".vue", ".svelte",
    # Java / JVM
    ".java", ".kt", ".scala", ".groovy", ".gradle",
    # C family
    ".c", ".cc", ".cxx", ".cpp", ".h", ".hpp", ".hxx",
    # .NET
    ".cs", ".fs", ".vb",
    # Systems / other
    ".go", ".rs", ".swift", ".dart", ".zig", ".nim",
    ".rb", ".php", ".pl", ".lua", ".r", ".m",
    ".ex", ".exs", ".clj", ".hs", ".elm", ".erl",
    # Data / config
    ".sql", ".graphql", ".proto",
    ".tf", ".hcl", ".dockerfile", ".makefile", ".cmake",
    ".gitignore", ".gitattributes",
}

_IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".tiff", ".tif", ".bmp", ".webp", ".gif", ".ico"}


def load_doc(filepath: str):
    """Load virtually any file type into LangChain Documents."""
    p   = Path(filepath)
    ext = p.suffix.lower()
    print(f"[VaultMind] Loading {p.name} (ext={ext})")

    try:
        # ── PDF ──────────────────────────────────────────
        if ext == ".pdf":
            from langchain_community.document_loaders import PyPDFLoader
            return PyPDFLoader(str(p)).load()

        # ── Word ─────────────────────────────────────────
        elif ext in (".docx", ".doc"):
            from langchain_community.document_loaders import Docx2txtLoader
            return Docx2txtLoader(str(p)).load()

        # ── PowerPoint ───────────────────────────────────
        elif ext in (".pptx", ".ppt"):
            return _load_pptx(str(p))

        # ── Excel ────────────────────────────────────────
        elif ext in (".xlsx", ".xls", ".xlsm"):
            return _load_excel(str(p))

        # ── CSV ──────────────────────────────────────────
        elif ext in (".csv", ".tsv"):
            from langchain_community.document_loaders.csv_loader import CSVLoader
            delim = "\t" if ext == ".tsv" else ","
            return CSVLoader(str(p), csv_args={"delimiter": delim}).load()

        # ── HTML ─────────────────────────────────────────
        elif ext in (".html", ".htm"):
            return _load_html(str(p))

        # ── JSON ─────────────────────────────────────────
        elif ext in (".json", ".jsonl", ".ndjson"):
            content = p.read_text(encoding="utf-8", errors="replace")
            from langchain.schema import Document
            return [Document(page_content=content, metadata={"source": str(p)})]

        # ── XML ──────────────────────────────────────────
        elif ext in (".xml", ".svg"):
            return _load_xml(str(p))

        # ── EPUB ─────────────────────────────────────────
        elif ext == ".epub":
            try:
                from langchain_community.document_loaders import UnstructuredEPubLoader
                return UnstructuredEPubLoader(str(p)).load()
            except Exception:
                return _load_fallback(str(p))

        # ── RTF ──────────────────────────────────────────
        elif ext == ".rtf":
            return _load_fallback(str(p))

        # ── Jupyter Notebook ─────────────────────────────
        elif ext == ".ipynb":
            return _load_notebook(str(p))

        # ── Email ────────────────────────────────────────
        elif ext in (".eml", ".msg"):
            return _load_fallback(str(p))

        # ── Images (OCR) ─────────────────────────────────
        elif ext in _IMAGE_EXTS:
            return _load_image(str(p))

        # ── Code / Plain text ────────────────────────────
        elif ext in _TEXT_EXTS or ext == "":
            from langchain_community.document_loaders import TextLoader
            return TextLoader(str(p), encoding="utf-8", autodetect_encoding=True).load()

        # ── Unknown: cascade fallbacks ───────────────────
        else:
            return _load_fallback(str(p))

    except Exception as e:
        print(f"[VaultMind] Error loading {filepath}: {e}")
        return _load_raw(str(p))


# ── Format-specific helpers ────────────────────────────────────────────────────

def _load_pptx(filepath: str):
    try:
        from pptx import Presentation
        prs = Presentation(filepath)
        slides = []
        for i, slide in enumerate(prs.slides, 1):
            texts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    texts.append(shape.text.strip())
            if texts:
                slides.append(f"[Slide {i}]\n" + "\n".join(texts))
        from langchain.schema import Document
        return [Document(
            page_content="\n\n".join(slides) or "[Empty presentation]",
            metadata={"source": filepath}
        )]
    except ImportError:
        return _load_fallback(filepath)
    except Exception as e:
        print(f"[VaultMind] PPTX error: {e}")
        return _load_fallback(filepath)


def _load_excel(filepath: str):
    try:
        import openpyxl
        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
        from langchain.schema import Document
        docs = []
        for sheet in wb.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    rows.append("\t".join(cells))
            if rows:
                docs.append(Document(
                    page_content=f"[Sheet: {sheet.title}]\n" + "\n".join(rows),
                    metadata={"source": filepath, "sheet": sheet.title}
                ))
        wb.close()
        return docs if docs else _load_fallback(filepath)
    except ImportError:
        return _load_fallback(filepath)
    except Exception as e:
        print(f"[VaultMind] Excel error: {e}")
        return _load_fallback(filepath)


def _load_html(filepath: str):
    try:
        from langchain_community.document_loaders import BSHTMLLoader
        return BSHTMLLoader(filepath).load()
    except Exception:
        try:
            from langchain_community.document_loaders import TextLoader
            return TextLoader(filepath, encoding="utf-8", autodetect_encoding=True).load()
        except Exception:
            return _load_raw(filepath)


def _load_xml(filepath: str):
    import re
    content = Path(filepath).read_text(encoding="utf-8", errors="replace")
    clean   = re.sub(r"<[^>]+>", " ", content)
    clean   = re.sub(r"\s+", " ", clean).strip()
    from langchain.schema import Document
    return [Document(page_content=clean or content, metadata={"source": filepath})]


def _load_notebook(filepath: str):
    try:
        import json as _json
        nb    = _json.loads(Path(filepath).read_text(encoding="utf-8"))
        cells = []
        for cell in nb.get("cells", []):
            ctype = cell.get("cell_type", "")
            src   = "".join(cell.get("source", []))
            if src.strip():
                cells.append(f"[{ctype.upper()}]\n{src}")
        from langchain.schema import Document
        return [Document(
            page_content="\n\n---\n\n".join(cells) or "[Empty notebook]",
            metadata={"source": filepath}
        )]
    except Exception as e:
        print(f"[VaultMind] Notebook error: {e}")
        return _load_raw(filepath)


def _load_image(filepath: str):
    from langchain.schema import Document
    name = Path(filepath).name
    try:
        import pytesseract
        from PIL import Image
        img  = Image.open(filepath)
        text = pytesseract.image_to_string(img)
        if text.strip():
            return [Document(page_content=text, metadata={"source": filepath, "type": "image_ocr"})]
        return [Document(page_content=f"[Image: {name} — OCR found no text]", metadata={"source": filepath})]
    except ImportError:
        # OCR not available — try to at least get image metadata
        try:
            from PIL import Image
            img = Image.open(filepath)
            return [Document(
                page_content=f"[Image: {name}, size: {img.size[0]}×{img.size[1]}, mode: {img.mode}. "
                             f"Install pytesseract for text extraction.]",
                metadata={"source": filepath}
            )]
        except Exception:
            return [Document(
                page_content=f"[Image: {name} — install Pillow + pytesseract for OCR]",
                metadata={"source": filepath}
            )]
    except Exception as e:
        return [Document(page_content=f"[Image: {name} — error: {e}]", metadata={"source": filepath})]


def _load_fallback(filepath: str):
    """Unstructured → TextLoader → raw bytes."""
    try:
        from langchain_community.document_loaders import UnstructuredFileLoader
        return UnstructuredFileLoader(filepath).load()
    except Exception:
        try:
            from langchain_community.document_loaders import TextLoader
            return TextLoader(filepath, encoding="utf-8", autodetect_encoding=True).load()
        except Exception:
            return _load_raw(filepath)


def _load_raw(filepath: str):
    """Absolute last resort — decode bytes as UTF-8 with error replacement."""
    try:
        content = Path(filepath).read_bytes().decode("utf-8", errors="replace")
        from langchain.schema import Document
        return [Document(page_content=content, metadata={"source": filepath})]
    except Exception:
        return []

def build_index(chroma_dir: str, file_paths: list) -> bool:
    """Rebuild a ChromaDB collection from a list of file paths."""
    from langchain.text_splitter import RecursiveCharacterTextSplitter
    from langchain_community.vectorstores import Chroma

    if Path(chroma_dir).exists():
        shutil.rmtree(chroma_dir)

    docs = []
    for fp in file_paths:
        docs.extend(load_doc(fp))

    if not docs:
        return False

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE, chunk_overlap=CHUNK_OVERLAP, add_start_index=True
    )
    chunks = splitter.split_documents(docs)
    Chroma.from_documents(chunks, get_emb(), persist_directory=chroma_dir)
    print(f"[VaultMind] Indexed {len(chunks)} chunks into {chroma_dir}")
    return True

def run_rag(chroma_dir: str, conv_id: int, query: str):
    """Run RAG chain, reconstructing memory from DB."""
    from langchain_community.vectorstores import Chroma
    from langchain_groq import ChatGroq
    from langchain.chains import ConversationalRetrievalChain
    from langchain.memory import ConversationBufferWindowMemory

    vs = Chroma(persist_directory=chroma_dir, embedding_function=get_emb())

    # Reconstruct conversation memory from DB (exclude current query)
    with db() as c:
        rows = c.execute(
            "SELECT role, content FROM messages WHERE conversation_id=? ORDER BY created_at",
            (conv_id,)
        ).fetchall()

    memory = ConversationBufferWindowMemory(
        memory_key="chat_history", return_messages=True, output_key="answer", k=6
    )
    i = 0
    while i < len(rows) - 1:  # -1 skips the current user query (last row)
        if rows[i]["role"] == "user" and rows[i + 1]["role"] == "assistant":
            memory.save_context({"input": rows[i]["content"]}, {"answer": rows[i + 1]["content"]})
            i += 2
        else:
            i += 1

    llm = ChatGroq(model=GROQ_MODEL, groq_api_key=GROQ_KEY, temperature=0.2, max_tokens=1024)
    chain = ConversationalRetrievalChain.from_llm(
        llm=llm,
        retriever=vs.as_retriever(search_type="mmr", search_kwargs={"k": TOP_K, "fetch_k": 10}),
        memory=memory,
        return_source_documents=True,
        verbose=False,
    )
    result = chain.invoke({"question": query})
    answer = result["answer"]

    seen, sources = set(), []
    for doc in result.get("source_documents", []):
        src  = doc.metadata.get("source", "unknown")
        page = doc.metadata.get("page", "")
        label = Path(src).name + (f" · p.{int(page)+1}" if page != "" else "")
        if label not in seen:
            seen.add(label)
            sources.append(label)

    confidence = "HIGH" if sources else "MED"
    return answer, sources, confidence

def run_llm_only(query: str, history: list):
    """Plain LLM chat (no retrieval) with conversation history."""
    from langchain_groq import ChatGroq
    from langchain.schema import HumanMessage, AIMessage, SystemMessage

    llm = ChatGroq(model=GROQ_MODEL, groq_api_key=GROQ_KEY, temperature=0.3, max_tokens=1024)
    msgs = [SystemMessage(content=(
        "You are VaultMind, a helpful AI assistant for document intelligence. "
        "No documents have been uploaded yet — answer from your own knowledge, clearly and concisely."
    ))]
    i = 0
    while i < len(history) - 1:
        if history[i]["role"] == "user" and history[i + 1]["role"] == "assistant":
            msgs.append(HumanMessage(content=history[i]["content"]))
            msgs.append(AIMessage(content=history[i + 1]["content"]))
            i += 2
        else:
            i += 1
    msgs.append(HumanMessage(content=query))

    resp = llm.invoke(msgs)
    return resp.content, [], ""

# ─── Projects ─────────────────────────────────────────────────────────────────

@app.route("/api/projects", methods=["GET"])
def list_projects():
    with db() as c:
        rows = c.execute("SELECT * FROM projects ORDER BY updated_at DESC").fetchall()
    return jsonify([dict(r) for r in rows])

@app.route("/api/projects", methods=["POST"])
def create_project():
    d = request.get_json()
    with db() as c:
        cur = c.execute(
            "INSERT INTO projects (name, description, instructions) VALUES (?,?,?)",
            (d.get("name", "New Project"), d.get("description", ""), d.get("instructions", ""))
        )
        row = c.execute("SELECT * FROM projects WHERE id=?", (cur.lastrowid,)).fetchone()
    return jsonify(dict(row)), 201

@app.route("/api/projects/<int:pid>", methods=["GET"])
def get_project(pid):
    with db() as c:
        row = c.execute("SELECT * FROM projects WHERE id=?", (pid,)).fetchone()
    return jsonify(dict(row)) if row else ("", 404)

@app.route("/api/projects/<int:pid>", methods=["PUT"])
def update_project(pid):
    d = request.get_json()
    with db() as c:
        c.execute(
            "UPDATE projects SET name=?, description=?, instructions=?, updated_at=datetime('now') WHERE id=?",
            (d.get("name"), d.get("description"), d.get("instructions"), pid)
        )
        row = c.execute("SELECT * FROM projects WHERE id=?", (pid,)).fetchone()
    return jsonify(dict(row))

@app.route("/api/projects/<int:pid>", methods=["DELETE"])
def delete_project(pid):
    with db() as c:
        c.execute("DELETE FROM projects WHERE id=?", (pid,))
    scope = os.path.join(CHROMA_BASE, f"project_{pid}")
    if Path(scope).exists():
        shutil.rmtree(scope)
    return jsonify({"ok": True})

# ─── Project Files ────────────────────────────────────────────────────────────

@app.route("/api/projects/<int:pid>/files", methods=["GET"])
def list_proj_files(pid):
    with db() as c:
        rows = c.execute(
            "SELECT * FROM project_files WHERE project_id=? ORDER BY created_at DESC", (pid,)
        ).fetchall()
    return jsonify([dict(r) for r in rows])

@app.route("/api/projects/<int:pid>/files", methods=["POST"])
def upload_proj_file(pid):
    f = request.files.get("file")
    if not f:
        return jsonify({"error": "No file"}), 400
    save_dir = os.path.join(UPLOADS_DIR, f"project_{pid}")
    os.makedirs(save_dir, exist_ok=True)
    name  = Path(f.filename).name
    fpath = os.path.join(save_dir, name)
    f.save(fpath)
    fsize = os.path.getsize(fpath)
    with db() as c:
        cur = c.execute(
            "INSERT INTO project_files (project_id, filename, filepath, file_size, indexed) VALUES (?,?,?,?,0)",
            (pid, name, fpath, fsize)
        )
        fid = cur.lastrowid
        all_paths = [r["filepath"] for r in c.execute(
            "SELECT filepath FROM project_files WHERE project_id=?", (pid,)
        ).fetchall()]
    scope = os.path.join(CHROMA_BASE, f"project_{pid}")
    def _index():
        ok = build_index(scope, all_paths)
        if ok:
            with db() as c2:
                c2.execute("UPDATE project_files SET indexed=1 WHERE project_id=?", (pid,))
    threading.Thread(target=_index, daemon=True).start()
    return jsonify({"id": fid, "filename": name, "file_size": fsize, "indexed": 0}), 201

@app.route("/api/projects/<int:pid>/files/<int:fid>", methods=["DELETE"])
def delete_proj_file(pid, fid):
    with db() as c:
        row = c.execute(
            "SELECT * FROM project_files WHERE id=? AND project_id=?", (fid, pid)
        ).fetchone()
        if row:
            c.execute("DELETE FROM project_files WHERE id=?", (fid,))
            try:
                os.remove(row["filepath"])
            except Exception:
                pass
            remaining = [r["filepath"] for r in c.execute(
                "SELECT filepath FROM project_files WHERE project_id=?", (pid,)
            ).fetchall()]
    if row:
        scope = os.path.join(CHROMA_BASE, f"project_{pid}")
        if remaining:
            threading.Thread(target=build_index, args=(scope, remaining), daemon=True).start()
        elif Path(scope).exists():
            shutil.rmtree(scope)
    return jsonify({"ok": True})

@app.route("/api/projects/<int:pid>/index-status", methods=["GET"])
def proj_index_status(pid):
    with db() as c:
        total   = c.execute("SELECT COUNT(*) FROM project_files WHERE project_id=?",       (pid,)).fetchone()[0]
        indexed = c.execute("SELECT COUNT(*) FROM project_files WHERE project_id=? AND indexed=1", (pid,)).fetchone()[0]
    return jsonify({"total": total, "indexed": indexed, "ready": total > 0 and total == indexed})

# ─── Conversations ────────────────────────────────────────────────────────────

@app.route("/api/conversations", methods=["GET"])
def list_convs():
    pid = request.args.get("project_id")
    with db() as c:
        if pid:
            rows = c.execute(
                "SELECT * FROM conversations WHERE project_id=? ORDER BY updated_at DESC", (pid,)
            ).fetchall()
        else:
            rows = c.execute(
                "SELECT * FROM conversations WHERE project_id IS NULL ORDER BY updated_at DESC"
            ).fetchall()
    return jsonify([dict(r) for r in rows])

@app.route("/api/conversations", methods=["POST"])
def create_conv():
    d = request.get_json()
    with db() as c:
        cur = c.execute(
            "INSERT INTO conversations (project_id, title) VALUES (?,?)",
            (d.get("project_id"), d.get("title", "New conversation"))
        )
        row = c.execute("SELECT * FROM conversations WHERE id=?", (cur.lastrowid,)).fetchone()
    return jsonify(dict(row)), 201

@app.route("/api/conversations/<int:cid>", methods=["PUT"])
def update_conv(cid):
    d = request.get_json()
    with db() as c:
        c.execute(
            "UPDATE conversations SET title=?, updated_at=datetime('now') WHERE id=?",
            (d.get("title"), cid)
        )
        row = c.execute("SELECT * FROM conversations WHERE id=?", (cid,)).fetchone()
    return jsonify(dict(row))

@app.route("/api/conversations/<int:cid>", methods=["DELETE"])
def delete_conv(cid):
    with db() as c:
        c.execute("DELETE FROM conversations WHERE id=?", (cid,))
    scope = os.path.join(CHROMA_BASE, f"conv_{cid}")
    if Path(scope).exists():
        shutil.rmtree(scope)
    return jsonify({"ok": True})

# ─── Messages ─────────────────────────────────────────────────────────────────

@app.route("/api/conversations/<int:cid>/messages", methods=["GET"])
def get_msgs(cid):
    with db() as c:
        rows = c.execute(
            "SELECT * FROM messages WHERE conversation_id=? ORDER BY created_at", (cid,)
        ).fetchall()
    result = []
    for r in rows:
        d = dict(r)
        d["sources"] = json.loads(d.get("sources") or "[]")
        result.append(d)
    return jsonify(result)

@app.route("/api/conversations/<int:cid>/messages", methods=["POST"])
def send_msg(cid):
    d     = request.get_json()
    query = (d.get("content") or "").strip()
    if not query:
        return jsonify({"error": "Empty message"}), 400

    if not GROQ_KEY:
        return jsonify({"error": "GROQ_API_KEY not set — add it to .env"}), 500

    with db() as c:
        conv = c.execute("SELECT * FROM conversations WHERE id=?", (cid,)).fetchone()
        if not conv:
            return jsonify({"error": "Conversation not found"}), 404
        pid = conv["project_id"]

        # First user message? → auto-title the conversation
        user_count = c.execute(
            "SELECT COUNT(*) FROM messages WHERE conversation_id=? AND role='user'", (cid,)
        ).fetchone()[0]
        is_first = user_count == 0

        # Save user message
        c.execute(
            "INSERT INTO messages (conversation_id, role, content) VALUES (?,?,?)",
            (cid, "user", query)
        )
        if is_first:
            title = query[:50] + ("…" if len(query) > 50 else "")
            c.execute(
                "UPDATE conversations SET title=?, updated_at=datetime('now') WHERE id=?", (title, cid)
            )
        else:
            c.execute("UPDATE conversations SET updated_at=datetime('now') WHERE id=?", (cid,))

        # Determine scope
        if pid:
            scope_dir  = os.path.join(CHROMA_BASE, f"project_{pid}")
            file_count = c.execute(
                "SELECT COUNT(*) FROM project_files WHERE project_id=? AND indexed=1", (pid,)
            ).fetchone()[0]
        else:
            scope_dir  = os.path.join(CHROMA_BASE, f"conv_{cid}")
            file_count = c.execute(
                "SELECT COUNT(*) FROM conversation_files WHERE conversation_id=? AND indexed=1", (cid,)
            ).fetchone()[0]

        # History for LLM-only fallback (exclude current query = last row)
        hist = [dict(r) for r in c.execute(
            "SELECT role, content FROM messages WHERE conversation_id=? ORDER BY created_at", (cid,)
        ).fetchall()]

    # Run RAG or plain LLM
    try:
        if file_count > 0 and Path(scope_dir).exists():
            answer, sources, confidence = run_rag(scope_dir, cid, query)
        else:
            answer, sources, confidence = run_llm_only(query, hist[:-1])
    except Exception as e:
        print(f"[VaultMind] AI error: {e}")
        answer, sources, confidence = f"Error: {e}", [], "LOW"

    # Save assistant reply
    with db() as c:
        c.execute(
            "INSERT INTO messages (conversation_id, role, content, sources, confidence) VALUES (?,?,?,?,?)",
            (cid, "assistant", answer, json.dumps(sources), confidence)
        )
        msg_row  = c.execute(
            "SELECT * FROM messages WHERE conversation_id=? AND role='assistant' ORDER BY created_at DESC LIMIT 1", (cid,)
        ).fetchone()
        conv_row = c.execute("SELECT * FROM conversations WHERE id=?", (cid,)).fetchone()

    return jsonify({
        "message": {
            "id": msg_row["id"],
            "role": "assistant",
            "content": answer,
            "sources": sources,
            "confidence": confidence,
            "created_at": msg_row["created_at"],
        },
        "conversation": dict(conv_row),
    })

# ─── Conversation Files ───────────────────────────────────────────────────────

@app.route("/api/conversations/<int:cid>/files", methods=["GET"])
def list_conv_files(cid):
    with db() as c:
        rows = c.execute(
            "SELECT * FROM conversation_files WHERE conversation_id=? ORDER BY created_at DESC", (cid,)
        ).fetchall()
    return jsonify([dict(r) for r in rows])

@app.route("/api/conversations/<int:cid>/files", methods=["POST"])
def upload_conv_file(cid):
    f = request.files.get("file")
    if not f:
        return jsonify({"error": "No file"}), 400
    save_dir = os.path.join(UPLOADS_DIR, f"conv_{cid}")
    os.makedirs(save_dir, exist_ok=True)
    name  = Path(f.filename).name
    fpath = os.path.join(save_dir, name)
    f.save(fpath)
    fsize = os.path.getsize(fpath)
    with db() as c:
        cur = c.execute(
            "INSERT INTO conversation_files (conversation_id, filename, filepath, file_size, indexed) VALUES (?,?,?,?,0)",
            (cid, name, fpath, fsize)
        )
        fid = cur.lastrowid
        all_paths = [r["filepath"] for r in c.execute(
            "SELECT filepath FROM conversation_files WHERE conversation_id=?", (cid,)
        ).fetchall()]
    scope = os.path.join(CHROMA_BASE, f"conv_{cid}")
    def _index():
        ok = build_index(scope, all_paths)
        if ok:
            with db() as c2:
                c2.execute("UPDATE conversation_files SET indexed=1 WHERE conversation_id=?", (cid,))
    threading.Thread(target=_index, daemon=True).start()
    return jsonify({"id": fid, "filename": name, "file_size": fsize, "indexed": 0}), 201

@app.route("/api/conversations/<int:cid>/files/<int:fid>", methods=["DELETE"])
def delete_conv_file(cid, fid):
    with db() as c:
        row = c.execute(
            "SELECT * FROM conversation_files WHERE id=? AND conversation_id=?", (fid, cid)
        ).fetchone()
        if row:
            c.execute("DELETE FROM conversation_files WHERE id=?", (fid,))
            try:
                os.remove(row["filepath"])
            except Exception:
                pass
            remaining = [r["filepath"] for r in c.execute(
                "SELECT filepath FROM conversation_files WHERE conversation_id=?", (cid,)
            ).fetchall()]
    if row:
        scope = os.path.join(CHROMA_BASE, f"conv_{cid}")
        if remaining:
            threading.Thread(target=build_index, args=(scope, remaining), daemon=True).start()
        elif Path(scope).exists():
            shutil.rmtree(scope)
    return jsonify({"ok": True})

@app.route("/api/conversations/<int:cid>/index-status", methods=["GET"])
def conv_index_status(cid):
    with db() as c:
        total   = c.execute("SELECT COUNT(*) FROM conversation_files WHERE conversation_id=?",         (cid,)).fetchone()[0]
        indexed = c.execute("SELECT COUNT(*) FROM conversation_files WHERE conversation_id=? AND indexed=1", (cid,)).fetchone()[0]
    return jsonify({"total": total, "indexed": indexed, "ready": total > 0 and total == indexed})

# ─── Static ───────────────────────────────────────────────────────────────────

@app.route("/")
def index():
    return send_from_directory("static", "index.html")

# ─── Boot ─────────────────────────────────────────────────────────────────────

if __name__ == "__main__":
    init_db()
    print("\n  ╔══════════════════════════════════╗")
    print("  ║   VaultMind v2 — RAG + Projects  ║")
    print("  ║   http://localhost:5000           ║")
    print("  ╚══════════════════════════════════╝\n")
    app.run(debug=True, host="0.0.0.0", port=5000, threaded=True)
